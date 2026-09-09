from __future__ import annotations

import io
import re
from typing import Any


STATE_CODES = {
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA", "HI", "ID",
    "IL", "IN", "IA", "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS",
    "MO", "MT", "NE", "NV", "NH", "NJ", "NM", "NY", "NC", "ND", "OH", "OK",
    "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV",
    "WI", "WY", "DC",
}

STATE_NAMES = {
    "alabama", "alaska", "arizona", "arkansas", "california", "colorado", "connecticut",
    "delaware", "florida", "georgia", "hawaii", "idaho", "illinois", "indiana", "iowa",
    "kansas", "kentucky", "louisiana", "maine", "maryland", "massachusetts", "michigan",
    "minnesota", "mississippi", "missouri", "montana", "nebraska", "nevada",
    "new hampshire", "new jersey", "new mexico", "new york", "north carolina",
    "north dakota", "ohio", "oklahoma", "oregon", "pennsylvania", "rhode island",
    "south carolina", "south dakota", "tennessee", "texas", "utah", "vermont", "virginia",
    "washington", "west virginia", "wisconsin", "wyoming", "district of columbia",
}

TITLE_WORDS = {
    "analyst", "manager", "coordinator", "specialist", "director", "engineer", "developer",
    "administrator", "consultant", "representative", "supervisor", "lead", "associate",
    "technician", "recruiter", "planner", "accountant", "auditor", "scientist", "architect",
    "designer", "strategist", "officer", "executive", "advisor", "adviser", "controller",
    "buyer", "scheduler", "programmer", "trainer", "estimator", "inspector", "operator",
}

SKILL_TERMS = [
    "excel", "microsoft excel", "power bi", "tableau", "sql", "python", "r", "snowflake",
    "databricks", "aws", "azure", "google cloud", "gcp", "salesforce", "dynamics 365",
    "dataverse", "power automate", "power apps", "power platform", "sap", "oracle", "workday",
    "jira", "servicenow", "sharepoint", "smartsheet", "looker", "alteryx", "arcgis",
    "project management", "program management", "process improvement", "continuous improvement",
    "lean", "six sigma", "agile", "scrum", "stakeholder management", "requirements gathering",
    "business analysis", "change management", "data analysis", "data analytics",
    "data visualization", "reporting", "forecasting", "budgeting", "financial analysis",
    "accounting", "procurement", "purchasing", "supply chain", "logistics", "warehouse",
    "manufacturing", "quality assurance", "quality control", "sales operations", "sales",
    "customer success", "customer service", "marketing", "recruiting", "human resources",
    "payroll", "benefits", "healthcare", "clinical", "insurance", "compliance",
    "legal operations", "operations", "workflow", "root cause analysis", "training",
    "team leadership", "people management", "vendor management", "contract management",
    "crm", "data governance", "automation", "security roles", "system administration",
]

TOOL_TERMS = {
    "excel", "microsoft excel", "power bi", "tableau", "sql", "python", "r", "snowflake",
    "databricks", "aws", "azure", "google cloud", "gcp", "salesforce", "dynamics 365",
    "dataverse", "power automate", "power apps", "power platform", "sap", "oracle", "workday",
    "jira", "servicenow", "sharepoint", "smartsheet", "looker", "alteryx", "arcgis",
}

FOCUS_TERMS = {
    "project management", "program management", "process improvement", "continuous improvement",
    "operations", "workflow", "stakeholder management", "requirements gathering",
    "business analysis", "change management", "data analysis", "data analytics",
    "data visualization", "reporting", "forecasting", "financial analysis", "procurement",
    "supply chain", "logistics", "manufacturing", "quality assurance", "quality control",
    "sales operations", "customer success", "marketing", "recruiting", "human resources",
    "healthcare", "insurance", "compliance", "vendor management", "contract management",
    "data governance", "automation", "system administration",
}

DOMAIN_RULES = {
    "HR / People Operations": ["human resources", "hr operations", "hr generalist", "recruiting", "payroll", "benefits", "people operations"],
    "Healthcare / Clinical": ["healthcare", "clinical", "hospital", "health system", "patient", "payer", "provider", "pharmaceutical", "pharma"],
    "Finance / Accounting": ["accounting", "financial analysis", "finance", "fp&a", "gaap", "audit", "controller"],
    "Insurance": ["insurance", "claims", "underwriting", "actuarial", "property and casualty"],
    "Legal / Compliance": ["legal operations", "legal", "paralegal", "regulatory compliance", "compliance"],
    "Supply Chain / Procurement": ["supply chain", "procurement", "purchasing", "logistics", "warehouse", "buyer"],
}

EDUCATION_WORDS = {
    "university", "college", "school", "bachelor", "master", "mba", "degree", "associate degree",
    "certification", "certified", "certificate",
}


def _clean_lines(text: str) -> list[str]:
    lines = []
    for raw in text.replace("\r", "\n").split("\n"):
        line = re.sub(r"\s+", " ", raw).strip(" \t|•·-")
        if line:
            lines.append(line)
    return lines


def extract_resume_text(data: bytes, filename: str) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == "docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        chunks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    chunks.append(" | ".join(cells))
        return "\n".join(chunks)
    if suffix in {"txt", "text"}:
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Please upload a PDF, DOCX, or TXT résumé.")


def _guess_name(lines: list[str]) -> str:
    for line in lines[:8]:
        lower = line.lower()
        if "@" in line or "http" in lower or any(ch.isdigit() for ch in line):
            continue
        words = re.findall(r"[A-Za-z][A-Za-z'’-]+", line)
        if 2 <= len(words) <= 4 and len(line) <= 60:
            if not any(word in lower for word in TITLE_WORDS | EDUCATION_WORDS):
                return " ".join(words)
    return ""


def _guess_location(lines: list[str]) -> tuple[str, str]:
    for line in lines[:15]:
        match = re.search(r"\b([A-Za-z .'-]{2,40}),\s*([A-Z]{2})\b", line)
        if match and match.group(2) in STATE_CODES:
            return match.group(1).strip(), match.group(2)
    for line in lines[:15]:
        lower = line.lower()
        for state in sorted(STATE_NAMES, key=len, reverse=True):
            marker = ", " + state
            if marker in lower:
                city = line[: lower.index(marker)].strip(" ,|-")
                if city and len(city) <= 45:
                    return city, state.title()
    return "", ""


def _looks_like_title(line: str) -> bool:
    lower = line.lower()
    if len(line) > 100 or len(line.split()) > 12:
        return False
    if any(word in lower for word in ["education", "experience", "skills", "summary", "objective", "certification"]):
        return False
    if "@" in line or "http" in lower:
        return False
    words = set(re.findall(r"[a-z]+", lower))
    return bool(words & TITLE_WORDS)


def _guess_titles(lines: list[str]) -> list[str]:
    results: list[str] = []
    seen: set[str] = set()
    for line in lines:
        if not _looks_like_title(line):
            continue
        cleaned = re.sub(r"\b(?:19|20)\d{2}\b.*$", "", line).strip(" |,-–—")
        cleaned = re.sub(r"\s{2,}", " ", cleaned)
        if not cleaned or len(cleaned) > 90:
            continue
        key = cleaned.lower()
        if key not in seen:
            seen.add(key)
            results.append(cleaned)
        if len(results) >= 8:
            break
    return results


def _find_terms(text: str, terms: list[str] | set[str]) -> list[str]:
    lower = re.sub(r"\s+", " ", text.lower())
    found = []
    for term in sorted(terms, key=len, reverse=True):
        pattern = r"(?<![a-z0-9])" + re.escape(term.lower()) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            canonical = "Google Cloud" if term == "gcp" else term
            if canonical.lower() not in {x.lower() for x in found}:
                found.append(canonical)
    return found


def _guess_domains(text: str) -> list[str]:
    lower = text.lower()
    results = []
    for domain, terms in DOMAIN_RULES.items():
        hits = sum(1 for term in terms if term in lower)
        if hits >= 2 or (hits == 1 and any(lower.count(term) >= 2 for term in terms)):
            results.append(domain)
    return results


def _guess_seniority(titles: list[str]) -> list[str]:
    joined = " ".join(titles).lower()
    preferences = []
    for term in ["senior", "lead", "manager", "supervisor", "director", "analyst", "specialist", "coordinator", "associate"]:
        if term in joined:
            preferences.append(term)
    return preferences[:5]


def _education_lines(lines: list[str]) -> list[str]:
    results = []
    for line in lines:
        lower = line.lower()
        if any(term in lower for term in EDUCATION_WORDS) and len(line) <= 140:
            if line not in results:
                results.append(line)
        if len(results) >= 5:
            break
    return results


def parse_resume_text(text: str) -> dict[str, Any]:
    if not text or len(text.strip()) < 40:
        raise ValueError("I couldn't find enough readable text in that résumé.")

    lines = _clean_lines(text)
    titles = _guess_titles(lines)
    skills = _find_terms(text, SKILL_TERMS)
    tools = [term for term in skills if term.lower() in {x.lower() for x in TOOL_TERMS}]
    focus = [term for term in skills if term.lower() in {x.lower() for x in FOCUS_TERMS}]
    city, state = _guess_location(lines)

    strong_skills = skills[:16]
    secondary_skills = skills[16:28]
    target_titles = titles[:5]

    profile = {
        "name": _guess_name(lines),
        "home_city": city,
        "home_state": state,
        "preferred_location_terms": [x.lower() for x in [city, state] if x],
        "target_titles": target_titles,
        "title_family_terms": target_titles,
        "strong_skills": strong_skills,
        "secondary_skills": secondary_skills,
        "focus_terms": focus[:12],
        "bonus_terms": tools[:12],
        "seniority_preferences": _guess_seniority(titles),
        "domain_strengths": _guess_domains(text),
    }

    return {
        "profile": {key: value for key, value in profile.items() if value not in ("", [], None)},
        "summary": {
            "titles": titles[:8],
            "skills": skills[:24],
            "domains": _guess_domains(text),
            "education": _education_lines(lines),
            "characters_read": len(text),
        },
        "text": text,
    }


def parse_resume_file(data: bytes, filename: str) -> dict[str, Any]:
    text = extract_resume_text(data, filename)
    return parse_resume_text(text)
