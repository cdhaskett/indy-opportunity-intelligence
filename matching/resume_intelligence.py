from __future__ import annotations

import io
import re
from copy import deepcopy
from typing import Any

from docx import Document
from pypdf import PdfReader

ROLE_WORDS = (
    "analyst", "manager", "coordinator", "specialist", "administrator",
    "consultant", "developer", "engineer", "director", "lead", "supervisor",
)

SKILL_CATALOG = [
    "excel", "power bi", "dax", "power query", "tableau", "looker",
    "sql", "python", "r", "snowflake", "databricks", "aws", "azure",
    "salesforce", "dynamics 365", "dataverse", "power automate", "power apps",
    "power platform", "sharepoint", "jira", "confluence", "servicenow",
    "sap", "netsuite", "workday", "epicor", "oracle", "crm", "erp",
    "data analysis", "data visualization", "reporting", "dashboard",
    "requirements gathering", "business analysis", "stakeholder management",
    "process improvement", "continuous improvement", "process mapping",
    "root cause", "workflow", "automation", "etl", "elt", "data governance",
    "project management", "change management", "agile", "scrum",
    "geospatial", "arcgis", "streamlit", "github",
]

IMPACT_PATTERNS = [
    re.compile(r"\b\d+(?:\.\d+)?%\b"),
    re.compile(r"\$\s?\d"),
    re.compile(r"\b(?:increased|decreased|reduced|improved|saved|grew|cut|eliminated|automated)\b", re.I),
]


def normalize(text: str | None) -> str:
    return re.sub(r"\s+", " ", (text or "").lower()).strip()


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Extract readable text from PDF, DOCX, or TXT resumes."""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if suffix == "docx":
        doc = Document(io.BytesIO(file_bytes))
        blocks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return "\n".join(blocks).strip()
    if suffix in {"txt", "md"}:
        return file_bytes.decode("utf-8", errors="ignore").strip()
    raise ValueError("Please upload a PDF, DOCX, or TXT resume.")


def _find_skills(text: str) -> list[str]:
    haystack = normalize(text)
    found = []
    for skill in SKILL_CATALOG:
        if re.search(rf"(?<![a-z0-9]){re.escape(skill)}(?![a-z0-9])", haystack):
            found.append(skill)
    return found


def _find_titles(text: str) -> list[str]:
    titles: list[str] = []
    for line in text.splitlines():
        cleaned = re.sub(r"\s+", " ", line).strip(" -•|\t")
        low = cleaned.lower()
        if not cleaned or len(cleaned) > 90:
            continue
        if any(re.search(rf"\b{word}\b", low) for word in ROLE_WORDS):
            if not re.search(r"\b(summary|skills|education|experience|professional)\b", low):
                cleaned = re.sub(r"\s*[|–—]\s*.*$", "", cleaned).strip()
                if 2 <= len(cleaned.split()) <= 10 and cleaned.lower() not in [x.lower() for x in titles]:
                    titles.append(cleaned)
        if len(titles) >= 8:
            break
    return titles


def infer_profile_from_resume(text: str, template: dict[str, Any]) -> dict[str, Any]:
    """Create an editable candidate profile from evidence actually present in the resume."""
    profile = deepcopy(template)
    skills = _find_skills(text)
    titles = _find_titles(text)

    profile["strong_skills"] = skills[:14]
    profile["secondary_skills"] = skills[14:28]
    if titles:
        profile["target_titles"] = titles[:6]
        family_terms = []
        for title in titles:
            for word in ROLE_WORDS:
                if re.search(rf"\b{word}\b", title.lower()) and word not in family_terms:
                    family_terms.append(word)
        profile["title_family_terms"] = family_terms or profile.get("title_family_terms", [])
        profile["seniority_preferences"] = [x for x in ["analyst", "specialist", "manager", "consultant"] if x in family_terms]

    profile["resume_inferred"] = True
    return profile


def resume_health(text: str) -> dict[str, Any]:
    low = text.lower()
    word_count = len(re.findall(r"\b\w+\b", text))
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    bulletish = sum(1 for x in lines if x.startswith(("-", "•", "*", "▪", "◦")))
    impact_hits = sum(1 for p in IMPACT_PATTERNS if p.search(text))
    skills = _find_skills(text)

    section_hits = sum(
        1 for section in ["experience", "education", "skills"]
        if re.search(rf"\b{section}\b", low)
    )

    score = 35
    score += min(20, len(skills) * 2)
    score += min(15, bulletish * 2)
    score += impact_hits * 8
    score += section_hits * 5
    if 250 <= word_count <= 900:
        score += 10
    score = max(0, min(100, score))

    notes = []
    if len(skills) < 6:
        notes.append("Skills are thin or not named explicitly. Add tools, systems, and methods you actually use.")
    if bulletish < 4:
        notes.append("The resume has few clear accomplishment bullets. Make experience easier to scan.")
    if impact_hits < 2:
        notes.append("Add measurable outcomes where possible: time saved, adoption, dollars, volume, accuracy, or percentage change.")
    if section_hits < 3:
        notes.append("Use clear Experience, Education, and Skills headings so both humans and ATS tools can parse it.")
    if word_count < 250:
        notes.append("The resume may be too sparse to prove the full background.")
    elif word_count > 1000:
        notes.append("The resume is very long. Tightening it may make the strongest evidence easier to find.")

    return {
        "score": score,
        "word_count": word_count,
        "skills_found": skills,
        "notes": notes,
    }


def resume_match_score(job: dict[str, Any], resume_text: str, profile: dict[str, Any]) -> dict[str, Any]:
    """Estimate how well the actual resume text proves the fit represented by the candidate profile."""
    job_text = normalize(f"{job.get('title', '')} {job.get('description', '')}")
    resume = normalize(resume_text)

    profile_skills = list(dict.fromkeys(profile.get("strong_skills", []) + profile.get("secondary_skills", [])))
    relevant_skills = [s for s in profile_skills if normalize(s) in job_text]
    proven_skills = [s for s in relevant_skills if normalize(s) in resume]

    target_titles = profile.get("target_titles", [])
    title_terms = [t for t in target_titles if normalize(t) in job_text]
    proven_titles = [t for t in title_terms if normalize(t) in resume]

    focus_terms = profile.get("focus_terms", [])
    relevant_focus = [t for t in focus_terms if normalize(t) in job_text]
    proven_focus = [t for t in relevant_focus if normalize(t) in resume]

    denom = max(1, len(relevant_skills) * 3 + len(title_terms) * 2 + len(relevant_focus))
    numer = len(proven_skills) * 3 + len(proven_titles) * 2 + len(proven_focus)
    score = round(100 * numer / denom)

    missing = [s for s in relevant_skills if s not in proven_skills]
    return {
        "score": max(0, min(100, score)),
        "proven_skills": proven_skills,
        "missing_skills": missing,
        "relevant_skills": relevant_skills,
    }
